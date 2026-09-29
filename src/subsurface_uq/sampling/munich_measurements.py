from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np

Array = np.ndarray

MISSING_IDENTIFIER_TOKENS = {"", "fehlt", "nicht bekannt", "none", "nan"}


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "Munich measurement loading requires openpyxl. Install with "
            "python -m pip install -e \".[geostat]\""
        ) from exc
    return openpyxl


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def hydraulic_conductivity_to_intrinsic_permeability(
    hydraulic_conductivity_m_s: Array,
    *,
    dynamic_viscosity_pa_s: float = 1.002e-3,
    density_kg_m3: float = 998.2,
    gravity_m_s2: float = 9.80665,
) -> Array:
    """Convert hydraulic conductivity K [m/s] to intrinsic permeability k [m^2].

    Defaults correspond approximately to liquid water near 20 degC. The chosen
    constants should be recorded with any downstream LGCNN experiment because
    they shift log10(K) by a constant when converted to log10(k).
    """

    values = np.asarray(hydraulic_conductivity_m_s, dtype=np.float64)
    if not np.all(np.isfinite(values)) or np.any(values <= 0.0):
        raise ValueError("hydraulic conductivity must be finite and positive")
    mu = float(dynamic_viscosity_pa_s)
    rho = float(density_kg_m3)
    g = float(gravity_m_s2)
    if mu <= 0.0 or rho <= 0.0 or g <= 0.0:
        raise ValueError("viscosity, density and gravity must be positive")
    return values * mu / (rho * g)


@dataclass(frozen=True)
class ReferenceHorizontalGrid:
    x_min_m: float
    y_min_m: float
    cell_size_x_m: float
    cell_size_y_m: float
    nx: int
    ny: int
    active_mask: Array

    @property
    def shape(self) -> tuple[int, int]:
        return (self.ny, self.nx)

    def nearest_indices(self, x_m: Array, y_m: Array) -> tuple[Array, Array, Array]:
        x = np.asarray(x_m, dtype=np.float64)
        y = np.asarray(y_m, dtype=np.float64)
        ix = np.rint((x - self.x_min_m) / self.cell_size_x_m).astype(np.int64)
        iy = np.rint((y - self.y_min_m) / self.cell_size_y_m).astype(np.int64)
        in_bounds = (ix >= 0) & (ix < self.nx) & (iy >= 0) & (iy < self.ny)
        active = np.zeros(ix.shape, dtype=bool)
        valid = np.flatnonzero(in_bounds)
        if valid.size:
            active[valid] = self.active_mask[iy[valid], ix[valid]]
        return ix, iy, in_bounds & active

    def center_coordinates(self, ix: Array, iy: Array) -> tuple[Array, Array]:
        ix = np.asarray(ix, dtype=np.int64)
        iy = np.asarray(iy, dtype=np.int64)
        return (
            self.x_min_m + ix * self.cell_size_x_m,
            self.y_min_m + iy * self.cell_size_y_m,
        )


def load_reference_horizontal_grid(
    path: str | Path,
    *,
    expected_cell_size_m: float | None = 100.0,
) -> ReferenceHorizontalGrid:
    """Scan the headerless Munich 3-D K table and recover its active XY footprint.

    The file is whitespace-delimited despite its .csv suffix. Only X and Y are
    used; repeated vertical levels are collapsed while streaming, so the full
    three-dimensional table is never retained in memory.
    """

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    xy: set[tuple[float, float]] = set()
    last: tuple[float, float] | None = None
    with source.open("r", encoding="utf-8", errors="replace") as handle:
        for line_number, line in enumerate(handle, start=1):
            parts = line.split()
            if not parts:
                continue
            if len(parts) < 2:
                raise ValueError(f"invalid reference-grid row {line_number}")
            try:
                pair = (float(parts[0]), float(parts[1]))
            except ValueError as exc:
                raise ValueError(
                    f"reference-grid row {line_number} does not begin with numeric X Y"
                ) from exc
            if pair != last:
                xy.add(pair)
                last = pair
    if not xy:
        raise ValueError("reference grid contains no XY coordinates")

    coordinates = np.asarray(sorted(xy), dtype=np.float64)
    x_values = np.unique(coordinates[:, 0])
    y_values = np.unique(coordinates[:, 1])
    if x_values.size < 2 or y_values.size < 2:
        raise ValueError("reference grid needs at least two X and Y positions")
    dx = float(np.median(np.diff(x_values)))
    dy = float(np.median(np.diff(y_values)))
    if dx <= 0.0 or dy <= 0.0:
        raise ValueError("reference-grid spacing must be positive")
    if expected_cell_size_m is not None:
        expected = float(expected_cell_size_m)
        if not np.isclose(dx, expected, rtol=0.0, atol=max(1e-6, expected * 1e-3)):
            raise ValueError(f"unexpected X spacing {dx:g} m; expected about {expected:g} m")
        if not np.isclose(dy, expected, rtol=0.0, atol=max(1e-6, expected * 1e-3)):
            raise ValueError(f"unexpected Y spacing {dy:g} m; expected about {expected:g} m")

    x_min = float(x_values[0])
    y_min = float(y_values[0])
    nx = int(x_values.size)
    ny = int(y_values.size)
    mask = np.zeros((ny, nx), dtype=bool)
    ix = np.rint((coordinates[:, 0] - x_min) / dx).astype(np.int64)
    iy = np.rint((coordinates[:, 1] - y_min) / dy).astype(np.int64)
    mask[iy, ix] = True
    return ReferenceHorizontalGrid(
        x_min_m=x_min,
        y_min_m=y_min,
        cell_size_x_m=dx,
        cell_size_y_m=dy,
        nx=nx,
        ny=ny,
        active_mask=mask,
    )


@dataclass(frozen=True)
class MunichHydraulicConductivityMeasurements:
    measurement_id: Array
    excel_row: Array
    x_m: Array
    y_m: Array
    hydraulic_conductivity_m_s: Array
    source: Array
    district: Array
    stratigraphy: Array
    groundwater_state: Array
    ix: Array | None = None
    iy: Array | None = None
    nearest_cell_displacement_m: Array | None = None

    def __len__(self) -> int:
        return int(self.hydraulic_conductivity_m_s.size)

    @property
    def log10_hydraulic_conductivity(self) -> Array:
        return np.log10(self.hydraulic_conductivity_m_s)

    @property
    def coordinates_xy_m(self) -> Array:
        return np.column_stack((self.x_m, self.y_m))

    def subset(self, mask: Array) -> "MunichHydraulicConductivityMeasurements":
        mask = np.asarray(mask, dtype=bool)
        if mask.shape != (len(self),):
            raise ValueError("subset mask has wrong shape")
        def optional(values):
            return None if values is None else values[mask]
        return MunichHydraulicConductivityMeasurements(
            measurement_id=self.measurement_id[mask],
            excel_row=self.excel_row[mask],
            x_m=self.x_m[mask],
            y_m=self.y_m[mask],
            hydraulic_conductivity_m_s=self.hydraulic_conductivity_m_s[mask],
            source=self.source[mask],
            district=self.district[mask],
            stratigraphy=self.stratigraphy[mask],
            groundwater_state=self.groundwater_state[mask],
            ix=optional(self.ix),
            iy=optional(self.iy),
            nearest_cell_displacement_m=optional(self.nearest_cell_displacement_m),
        )

    def to_rows(self) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for index in range(len(self)):
            row: dict[str, object] = {
                "measurement_id": str(self.measurement_id[index]),
                "excel_row": int(self.excel_row[index]),
                "x_m": float(self.x_m[index]),
                "y_m": float(self.y_m[index]),
                "hydraulic_conductivity_m_s": float(self.hydraulic_conductivity_m_s[index]),
                "log10_hydraulic_conductivity": float(self.log10_hydraulic_conductivity[index]),
                "source": str(self.source[index]),
                "district": str(self.district[index]),
                "stratigraphy": str(self.stratigraphy[index]),
                "groundwater_state": str(self.groundwater_state[index]),
            }
            if self.ix is not None:
                row["ix"] = int(self.ix[index])
                row["iy"] = int(self.iy[index])
                row["nearest_cell_displacement_m"] = float(
                    self.nearest_cell_displacement_m[index]
                )
            rows.append(row)
        return rows


def load_munich_hydraulic_conductivity_measurements(
    path: str | Path,
    *,
    sheet_name: str = "kf_werte_180223",
    stratigraphy: str = "q",
    groundwater_state: str = "ungespannt",
    reference_grid: ReferenceHorizontalGrid | None = None,
) -> tuple[MunichHydraulicConductivityMeasurements, dict[str, object]]:
    """Load, QC and filter the real Munich hydraulic-conductivity observations."""

    openpyxl = _require_openpyxl()
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"sheet {sheet_name!r} not found; available: {workbook.sheetnames}")
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    header = next(iterator)
    columns = {_clean_text(value): index for index, value in enumerate(header)}
    required = [
        "Ob_Id", "Rechtswert", "Hochwert", "BEARBEITER", "KF Wert",
        "Landkreis", "Strategrap", "GW_Zustand",
    ]
    missing = [name for name in required if name not in columns]
    if missing:
        raise ValueError(f"measurement workbook is missing columns: {missing}")

    accepted: list[dict[str, object]] = []
    counters = {
        "total_rows": 0,
        "non_numeric_or_missing_k": 0,
        "non_positive_k": 0,
        "non_numeric_coordinates": 0,
        "stratigraphy_filtered": 0,
        "groundwater_state_filtered": 0,
        "outside_active_footprint": 0,
    }
    for excel_row, row in enumerate(iterator, start=2):
        counters["total_rows"] += 1
        k_raw = row[columns["KF Wert"]]
        if not isinstance(k_raw, (int, float)) or isinstance(k_raw, bool) or not np.isfinite(k_raw):
            counters["non_numeric_or_missing_k"] += 1
            continue
        k_value = float(k_raw)
        if k_value <= 0.0:
            counters["non_positive_k"] += 1
            continue
        x_raw = row[columns["Rechtswert"]]
        y_raw = row[columns["Hochwert"]]
        if (
            not isinstance(x_raw, (int, float))
            or not isinstance(y_raw, (int, float))
            or isinstance(x_raw, bool)
            or isinstance(y_raw, bool)
            or not np.isfinite(x_raw)
            or not np.isfinite(y_raw)
        ):
            counters["non_numeric_coordinates"] += 1
            continue
        strat = _clean_text(row[columns["Strategrap"]])
        gw = _clean_text(row[columns["GW_Zustand"]])
        if strat.lower() != stratigraphy.lower():
            counters["stratigraphy_filtered"] += 1
            continue
        if gw.lower() != groundwater_state.lower():
            counters["groundwater_state_filtered"] += 1
            continue
        object_id = _clean_text(row[columns["Ob_Id"]])
        if object_id.lower() in MISSING_IDENTIFIER_TOKENS:
            object_id = f"excel_row_{excel_row}"
        accepted.append(
            {
                "measurement_id": object_id,
                "excel_row": excel_row,
                "x_m": float(x_raw),
                "y_m": float(y_raw),
                "k": k_value,
                "source": _clean_text(row[columns["BEARBEITER"]]),
                "district": _clean_text(row[columns["Landkreis"]]),
                "stratigraphy": strat,
                "groundwater_state": gw,
            }
        )
    workbook.close()

    if not accepted:
        raise ValueError("no measurements survive workbook QC and semantic filters")
    x = np.asarray([row["x_m"] for row in accepted], dtype=np.float64)
    y = np.asarray([row["y_m"] for row in accepted], dtype=np.float64)
    ix = iy = displacement = None
    if reference_grid is not None:
        mapped_ix, mapped_iy, active = reference_grid.nearest_indices(x, y)
        counters["outside_active_footprint"] = int(np.count_nonzero(~active))
        accepted = [row for row, keep in zip(accepted, active.tolist()) if keep]
        if not accepted:
            raise ValueError("no measurements lie on the active reference-grid footprint")
        x = x[active]
        y = y[active]
        ix = mapped_ix[active]
        iy = mapped_iy[active]
        center_x, center_y = reference_grid.center_coordinates(ix, iy)
        displacement = np.sqrt((x - center_x) ** 2 + (y - center_y) ** 2)

    measurements = MunichHydraulicConductivityMeasurements(
        measurement_id=np.asarray([row["measurement_id"] for row in accepted], dtype=object),
        excel_row=np.asarray([row["excel_row"] for row in accepted], dtype=np.int64),
        x_m=x,
        y_m=y,
        hydraulic_conductivity_m_s=np.asarray([row["k"] for row in accepted], dtype=np.float64),
        source=np.asarray([row["source"] for row in accepted], dtype=object),
        district=np.asarray([row["district"] for row in accepted], dtype=object),
        stratigraphy=np.asarray([row["stratigraphy"] for row in accepted], dtype=object),
        groundwater_state=np.asarray([row["groundwater_state"] for row in accepted], dtype=object),
        ix=ix,
        iy=iy,
        nearest_cell_displacement_m=displacement,
    )

    log_values = measurements.log10_hydraulic_conductivity
    metadata: dict[str, object] = {
        **counters,
        "accepted_measurements": len(measurements),
        "sheet_name": sheet_name,
        "stratigraphy_filter": stratigraphy,
        "groundwater_state_filter": groundwater_state,
        "coordinate_reference_system_assumption": "DHDN / Gauss-Kruger zone 4 (EPSG:31468), pending confirmation",
        "hydraulic_conductivity_units": "m/s",
        "min_k_m_s": float(np.min(measurements.hydraulic_conductivity_m_s)),
        "median_k_m_s": float(np.median(measurements.hydraulic_conductivity_m_s)),
        "mean_k_m_s": float(np.mean(measurements.hydraulic_conductivity_m_s)),
        "max_k_m_s": float(np.max(measurements.hydraulic_conductivity_m_s)),
        "q05_k_m_s": float(np.quantile(measurements.hydraulic_conductivity_m_s, 0.05)),
        "q95_k_m_s": float(np.quantile(measurements.hydraulic_conductivity_m_s, 0.95)),
        "mean_log10_k": float(np.mean(log_values)),
        "std_log10_k": float(np.std(log_values, ddof=1)),
        "above_nominal_5e-2_count": int(np.count_nonzero(measurements.hydraulic_conductivity_m_s > 5.0e-2)),
    }
    if displacement is not None:
        metadata.update(
            {
                "distinct_reference_cells": int(
                    len(set(zip(measurements.iy.tolist(), measurements.ix.tolist())))
                ),
                "median_nearest_cell_displacement_m": float(np.median(displacement)),
                "max_nearest_cell_displacement_m": float(np.max(displacement)),
            }
        )
    return measurements, metadata


def aggregate_measurements_by_reference_cell(
    measurements: MunichHydraulicConductivityMeasurements,
) -> list[dict[str, object]]:
    """Return one geometric-mean K value per occupied reference cell."""

    if measurements.ix is None or measurements.iy is None:
        raise ValueError("measurements must first be mapped to a reference grid")
    groups: dict[tuple[int, int], list[int]] = {}
    for index, cell in enumerate(zip(measurements.iy.tolist(), measurements.ix.tolist())):
        groups.setdefault(cell, []).append(index)
    rows: list[dict[str, object]] = []
    log_values = measurements.log10_hydraulic_conductivity
    for (iy, ix), indices in sorted(groups.items()):
        selected = log_values[np.asarray(indices, dtype=np.int64)]
        mean_log = float(np.mean(selected))
        rows.append(
            {
                "iy": int(iy),
                "ix": int(ix),
                "measurement_count": len(indices),
                "mean_log10_hydraulic_conductivity": mean_log,
                "geometric_mean_hydraulic_conductivity_m_s": float(10.0**mean_log),
                "within_cell_log10_variance": float(np.var(selected, ddof=1)) if len(indices) > 1 else 0.0,
            }
        )
    return rows
